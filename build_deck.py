import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

def create_deck(output_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette - Modern Deep Obsidian / Tech Glass Theme
    C_BG = RGBColor(11, 15, 25)          # #0B0F19
    C_CARD_BG = RGBColor(18, 26, 44)     # #121A2C
    C_CARD_ALT = RGBColor(24, 34, 56)    # #182238
    C_CARD_BORDER = RGBColor(38, 51, 78) # #26334E
    
    C_TEXT_WHITE = RGBColor(255, 255, 255)
    C_TEXT_SLATE = RGBColor(203, 213, 225) # #CBD5E1
    C_TEXT_MUTED = RGBColor(148, 163, 184) # #94A3B8
    
    C_CYAN = RGBColor(56, 189, 248)       # #38BDF8 (Sky/Cyan)
    C_BLUE = RGBColor(59, 130, 246)       # #3B82F6 (Electric Blue)
    C_INDIGO = RGBColor(99, 102, 241)     # #6366F1
    C_EMERALD = RGBColor(52, 211, 153)    # #34D399 (Emerald Green)
    C_AMBER = RGBColor(251, 146, 60)      # #FB923C (Amber/Orange)
    C_ROSE = RGBColor(244, 63, 94)        # #F43F5E

    FONT_HEADING = "Segoe UI"
    FONT_BODY = "Segoe UI"

    def set_bg(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = C_BG

        # Subtle decorative glow line at top
        top_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.04))
        top_line.fill.solid()
        top_line.fill.fore_color.rgb = C_BLUE
        top_line.line.fill.background()

    def add_card(slide, left, top, width, height, bg_color=C_CARD_BG, border_color=C_CARD_BORDER, border_width=Pt(1)):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = border_width
        else:
            card.line.fill.background()
        return card

    def add_badge(slide, left, top, width, height, text, bg_color, text_color, font_size=Pt(9)):
        badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        badge.fill.solid()
        badge.fill.fore_color.rgb = bg_color
        badge.line.fill.background()
        tf = badge.text_frame
        tf.clear()
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = text
        run.font.name = FONT_HEADING
        run.font.size = font_size
        run.font.bold = True
        run.font.color.rgb = text_color
        return badge

    def add_slide_header(slide, num_str, tag_str, title_str, subtitle_str=None):
        # Tag Badge Top Left
        add_badge(slide, Inches(0.8), Inches(0.45), Inches(2.2), Inches(0.28), tag_str, RGBColor(20, 38, 68), C_CYAN, Pt(9))
        
        # Slide Counter Top Right
        counter = slide.shapes.add_textbox(Inches(11.0), Inches(0.42), Inches(1.533), Inches(0.35))
        tf = counter.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        r1 = p.add_run()
        r1.text = num_str
        r1.font.name = FONT_HEADING
        r1.font.size = Pt(13)
        r1.font.bold = True
        r1.font.color.rgb = C_CYAN
        r2 = p.add_run()
        r2.text = "  /  05"
        r2.font.name = FONT_HEADING
        r2.font.size = Pt(11)
        r2.font.color.rgb = C_TEXT_MUTED

        # Main Title Box
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.733), Inches(0.75))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = Inches(0)
        tf_t.margin_top = Inches(0)
        p_t = tf_t.paragraphs[0]
        r_t = p_t.add_run()
        r_t.text = title_str
        r_t.font.name = FONT_HEADING
        r_t.font.size = Pt(22)
        r_t.font.bold = True
        r_t.font.color.rgb = C_TEXT_WHITE

        if subtitle_str:
            p_sub = tf_t.add_paragraph()
            r_sub = p_sub.add_run()
            r_sub.text = subtitle_str
            r_sub.font.name = FONT_BODY
            r_sub.font.size = Pt(12)
            r_sub.font.color.rgb = C_TEXT_SLATE

    # ==========================================
    # SLIDE 1: HERO / TITLE SLIDE
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_bg(slide1)

    # Top Brand Bar
    add_badge(slide1, Inches(0.8), Inches(0.6), Inches(1.8), Inches(0.32), "PROJECT CELIO", RGBColor(20, 38, 68), C_CYAN, Pt(9.5))
    
    tag_counter = slide1.shapes.add_textbox(Inches(11.0), Inches(0.55), Inches(1.533), Inches(0.35))
    tf = tag_counter.text_frame
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = "01 / 05"
    r.font.name = FONT_HEADING
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = C_TEXT_MUTED

    # Hero Main Card
    hero_card = add_card(slide1, Inches(0.8), Inches(1.2), Inches(11.733), Inches(3.4), bg_color=C_CARD_BG, border_color=C_CARD_BORDER)
    
    # Title Text Frame Inside Hero Card
    hero_tf = hero_card.text_frame
    hero_tf.word_wrap = True
    hero_tf.margin_left = Inches(0.5)
    hero_tf.margin_right = Inches(0.5)
    hero_tf.margin_top = Inches(0.4)

    # App Tag
    p_app = hero_tf.paragraphs[0]
    r_app = p_app.add_run()
    r_app.text = "COMPUTER VISION & AUTOMATION"
    r_app.font.name = FONT_HEADING
    r_app.font.size = Pt(11)
    r_app.font.bold = True
    r_app.font.color.rgb = C_CYAN

    # Main Project Title
    p_title = hero_tf.add_paragraph()
    p_title.space_before = Pt(8)
    r_title = p_title.add_run()
    r_title.text = "IsiDetector"
    r_title.font.name = FONT_HEADING
    r_title.font.size = Pt(44)
    r_title.font.bold = True
    r_title.font.color.rgb = C_TEXT_WHITE

    # Subtitle
    p_sub = hero_tf.add_paragraph()
    p_sub.space_before = Pt(6)
    r_sub = p_sub.add_run()
    r_sub.text = "Automated Parcel Sorting by Camera"
    r_sub.font.name = FONT_BODY
    r_sub.font.size = Pt(19)
    r_sub.font.color.rgb = C_TEXT_SLATE

    # Presenter Pill Card inside hero
    add_card(slide1, Inches(1.3), Inches(3.6), Inches(4.5), Inches(0.7), bg_color=C_CARD_ALT, border_color=RGBColor(48, 64, 96))
    pres_box = slide1.shapes.add_textbox(Inches(1.5), Inches(3.68), Inches(4.2), Inches(0.55))
    tf_p = pres_box.text_frame
    tf_p.word_wrap = True
    p1 = tf_p.paragraphs[0]
    r1 = p1.add_run()
    r1.text = "Atanda Abdullahi"
    r1.font.name = FONT_HEADING
    r1.font.size = Pt(13)
    r1.font.bold = True
    r1.font.color.rgb = C_TEXT_WHITE
    p2 = tf_p.add_paragraph()
    r2 = p2.add_run()
    r2.text = "Computer Vision & Edge Systems"
    r2.font.name = FONT_BODY
    r2.font.size = Pt(10)
    r2.font.color.rgb = C_CYAN

    # 4 Agenda Bento Cards at Bottom
    agenda_items = [
        ("01", "WHY", "The Problem & Objective", C_AMBER),
        ("02", "HOW", "3-Step Vision Pipeline", C_CYAN),
        ("03", "WHAT", "Results, Accuracy & Speed", C_EMERALD),
        ("04", "DEMO", "Live Sorter PLC Stream", C_INDIGO),
    ]

    card_w = Inches(2.78)
    card_gap = Inches(0.2)
    start_x = Inches(0.8)
    y_pos = Inches(4.85)
    card_h = Inches(2.0)

    for i, (num, label, desc, color) in enumerate(agenda_items):
        cx = start_x + i * (card_w + card_gap)
        acard = add_card(slide1, cx, y_pos, card_w, card_h, bg_color=C_CARD_BG, border_color=C_CARD_BORDER)
        
        # Step Number Badge
        add_badge(slide1, cx + Inches(0.25), y_pos + Inches(0.25), Inches(0.75), Inches(0.28), f"PART {num}", RGBColor(20, 38, 68), color, Pt(8.5))
        
        # Text box for label & desc
        tbox = slide1.shapes.add_textbox(cx + Inches(0.25), y_pos + Inches(0.65), card_w - Inches(0.5), Inches(1.1))
        tf_a = tbox.text_frame
        tf_a.word_wrap = True
        tf_a.margin_left = Inches(0)
        tf_a.margin_top = Inches(0)
        
        p_l = tf_a.paragraphs[0]
        r_l = p_l.add_run()
        r_l.text = label
        r_l.font.name = FONT_HEADING
        r_l.font.size = Pt(16)
        r_l.font.bold = True
        r_l.font.color.rgb = C_TEXT_WHITE
        
        p_d = tf_a.add_paragraph()
        p_d.space_before = Pt(4)
        r_d = p_d.add_run()
        r_d.text = desc
        r_d.font.name = FONT_BODY
        r_d.font.size = Pt(11)
        r_d.font.color.rgb = C_TEXT_SLATE

    # ==========================================
    # SLIDE 2: WHY (The Challenge & Solution)
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_bg(slide2)
    add_slide_header(slide2, "02", "THE PROBLEM & GOAL", "WHY IsiDetector?", "Eliminating manual sorting bottlenecks with real-time computer vision on high-speed conveyors.")

    # Two Column Layout
    col_w = Inches(5.72)
    col1_x = Inches(0.8)
    col2_x = Inches(6.81)
    
    # Left Column - 2 Problem Cards
    card1_h = Inches(2.55)
    card1 = add_card(slide2, col1_x, Inches(1.75), col_w, card1_h, bg_color=C_CARD_BG, border_color=C_CARD_BORDER)
    add_badge(slide2, col1_x + Inches(0.35), Inches(1.95), Inches(1.8), Inches(0.28), "MANUAL BOTTLENECK", RGBColor(50, 30, 20), C_AMBER, Pt(8.5))
    
    tb1 = slide2.shapes.add_textbox(col1_x + Inches(0.35), Inches(2.35), col_w - Inches(0.7), Inches(1.8))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = Inches(0)
    
    p1 = tf1.paragraphs[0]
    r1 = p1.add_run()
    r1.text = "Cartons vs. Polybags"
    r1.font.name = FONT_HEADING
    r1.font.size = Pt(16)
    r1.font.bold = True
    r1.font.color.rgb = C_TEXT_WHITE
    
    p1_b = tf1.add_paragraph()
    p1_b.space_before = Pt(6)
    r1_b = p1_b.add_run()
    r1_b.text = "Sorted by hand today in high-volume hubs. Manual sorting is slow, fatigue-prone, expensive, and inconsistent during peak throughput surges."
    r1_b.font.name = FONT_BODY
    r1_b.font.size = Pt(12)
    r1_b.font.color.rgb = C_TEXT_SLATE

    # Problem Card 2: Environment
    card2_y = Inches(4.5)
    card2 = add_card(slide2, col1_x, card2_y, col_w, card1_h, bg_color=C_CARD_BG, border_color=C_CARD_BORDER)
    add_badge(slide2, col1_x + Inches(0.35), card2_y + Inches(0.2), Inches(1.8), Inches(0.28), "HARSH ENVIRONMENT", RGBColor(50, 30, 20), C_AMBER, Pt(8.5))
    
    tb2 = slide2.shapes.add_textbox(col1_x + Inches(0.35), card2_y + Inches(0.6), col_w - Inches(0.7), Inches(1.8))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = Inches(0)
    
    p2 = tf2.paragraphs[0]
    r2 = p2.add_run()
    r2.text = "1 m/s Conveyor Dynamics"
    r2.font.name = FONT_HEADING
    r2.font.size = Pt(16)
    r2.font.bold = True
    r2.font.color.rgb = C_TEXT_WHITE
    
    p2_b = tf2.add_paragraph()
    p2_b.space_before = Pt(6)
    r2_b = p2_b.add_run()
    r2_b.text = "Parcels travel at high velocity under harsh warehouse lighting with severe plastic specular glare, dynamic shadows, and random orientations."
    r2_b.font.name = FONT_BODY
    r2_b.font.size = Pt(12)
    r2_b.font.color.rgb = C_TEXT_SLATE

    # Right Column - The Breakthrough Solution Hero Card
    hero_w = col_w
    hero_h = Inches(5.3)
    sol_card = add_card(slide2, col2_x, Inches(1.75), hero_w, hero_h, bg_color=RGBColor(16, 28, 52), border_color=C_BLUE, border_width=Pt(1.5))
    
    add_badge(slide2, col2_x + Inches(0.4), Inches(2.05), Inches(1.9), Inches(0.3), "THE BREAKTHROUGH", RGBColor(20, 50, 80), C_CYAN, Pt(9))
    
    sol_tb = slide2.shapes.add_textbox(col2_x + Inches(0.4), Inches(2.55), hero_w - Inches(0.8), Inches(4.3))
    sol_tf = sol_tb.text_frame
    sol_tf.word_wrap = True
    sol_tf.margin_left = Inches(0)
    
    p_s1 = sol_tf.paragraphs[0]
    r_s1 = p_s1.add_run()
    r_s1.text = "→ The Camera Decides."
    r_s1.font.name = FONT_HEADING
    r_s1.font.size = Pt(26)
    r_s1.font.bold = True
    r_s1.font.color.rgb = C_CYAN
    
    p_s2 = sol_tf.add_paragraph()
    p_s2.space_before = Pt(4)
    r_s2 = p_s2.add_run()
    r_s2.text = "Every parcel. In real time."
    r_s2.font.name = FONT_HEADING
    r_s2.font.size = Pt(17)
    r_s2.font.bold = True
    r_s2.font.color.rgb = C_TEXT_WHITE

    # Key Solution Pillars inside Card
    sol_points = [
        ("Autonomous Edge Vision", "Zero human intervention required. Edge model infers parcel class instantly on the fly."),
        ("Universal Hardware Support", "Deploys on existing site PCs (GPU or CPU Intel boxes with OpenVINO) + standard IP cameras."),
        ("Sub-millisecond Sorter Trigger", "Issues precise UDP datagrams to the PLC gate before the parcel reaches the divert zone."),
    ]

    for title, desc in sol_points:
        p_pt = sol_tf.add_paragraph()
        p_pt.space_before = Pt(14)
        r_pt_title = p_pt.add_run()
        r_pt_title.text = f"•  {title}: "
        r_pt_title.font.name = FONT_HEADING
        r_pt_title.font.size = Pt(12)
        r_pt_title.font.bold = True
        r_pt_title.font.color.rgb = C_TEXT_WHITE
        
        r_pt_desc = p_pt.add_run()
        r_pt_desc.text = desc
        r_pt_desc.font.name = FONT_BODY
        r_pt_desc.font.size = Pt(11.5)
        r_pt_desc.font.color.rgb = C_TEXT_SLATE

    # ==========================================
    # SLIDE 3: HOW (The 3-Step Vision Pipeline)
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_bg(slide3)
    add_slide_header(slide3, "03", "SYSTEM ARCHITECTURE", "HOW It Works: The 3-Step Pipeline", "From optical camera stream to real-time machine PLC sorter actuation.")

    left_w = Inches(6.5)
    right_x = Inches(7.55)
    right_w = Inches(4.98)

    # 3 Pipeline Step Cards on Left
    steps = [
        ("01", "SEE", "A camera watches the belt", "Captures live RTSP stream from any overhead IP camera. Operates seamlessly on CPU or GPU hardware without custom sensors.", C_CYAN),
        ("02", "DECIDE", "AI recognises & follows each parcel", "Real-time YOLO26 instance segmentation (polygon masks) + ByteTrack tracker. Trained on ~3,000 site images.", C_INDIGO),
        ("03", "TELL", "One message per parcel to the sorter", "Dispatches ultra-low-latency UDP datagram to PLC controller at virtual line crossing. Carton vs Polybag · counted once.", C_EMERALD),
    ]

    step_h = Inches(1.6)
    step_gap = Inches(0.18)
    step_start_y = Inches(1.75)

    for i, (s_num, s_badge, s_title, s_desc, s_color) in enumerate(steps):
        sy = step_start_y + i * (step_h + step_gap)
        scard = add_card(slide3, Inches(0.8), sy, left_w, step_h, bg_color=C_CARD_BG, border_color=C_CARD_BORDER)
        
        # Step Badge
        add_badge(slide3, Inches(1.05), sy + Inches(0.2), Inches(1.2), Inches(0.26), f"{s_num} · {s_badge}", RGBColor(20, 38, 68), s_color, Pt(8.5))
        
        # Text Frame
        tb_s = slide3.shapes.add_textbox(Inches(2.4), sy + Inches(0.15), left_w - Inches(1.8), step_h - Inches(0.3))
        tf_s = tb_s.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = Inches(0)
        tf_s.margin_top = Inches(0)
        
        p_st = tf_s.paragraphs[0]
        r_st = p_st.add_run()
        r_st.text = s_title
        r_st.font.name = FONT_HEADING
        r_st.font.size = Pt(13)
        r_st.font.bold = True
        r_st.font.color.rgb = C_TEXT_WHITE
        
        p_sd = tf_s.add_paragraph()
        p_sd.space_before = Pt(3)
        r_sd = p_sd.add_run()
        r_sd.text = s_desc
        r_sd.font.name = FONT_BODY
        r_sd.font.size = Pt(10.5)
        r_sd.font.color.rgb = C_TEXT_SLATE

    # Right Side: Visual Detection Showcase Card with image2.png
    vis_card = add_card(slide3, right_x, Inches(1.75), right_w, Inches(5.16), bg_color=C_CARD_BG, border_color=C_CARD_BORDER)
    
    # Image Frame Title
    add_badge(slide3, right_x + Inches(0.3), Inches(1.95), Inches(2.2), Inches(0.26), "LIVE DETECTION FEED", RGBColor(20, 38, 68), C_CYAN, Pt(8.5))
    
    # Add image2.png (detection image)
    img_path = "/home/aatanda/logistic/.pptx_media/image2.png"
    if os.path.exists(img_path):
        img_left = right_x + Inches(0.3)
        img_top = Inches(2.35)
        img_width = right_w - Inches(0.6)
        img_height = Inches(3.2)
        slide3.shapes.add_picture(img_path, img_left, img_top, img_width, img_height)

    # Detection Caption Box
    cap_box = slide3.shapes.add_textbox(right_x + Inches(0.3), Inches(5.7), right_w - Inches(0.6), Inches(1.0))
    tf_c = cap_box.text_frame
    tf_c.word_wrap = True
    tf_c.margin_left = Inches(0)
    
    p_c1 = tf_c.paragraphs[0]
    r_c1 = p_c1.add_run()
    r_c1.text = "Instance Segmentation & Virtual Line"
    r_c1.font.name = FONT_HEADING
    r_c1.font.size = Pt(11.5)
    r_c1.font.bold = True
    r_c1.font.color.rgb = C_TEXT_WHITE
    
    p_c2 = tf_c.add_paragraph()
    r_c2 = p_c2.add_run()
    r_c2.text = "Mask mAP 94.2% · Sub-15ms inference · Zero double-counting"
    r_c2.font.name = FONT_BODY
    r_c2.font.size = Pt(10)
    r_c2.font.color.rgb = C_CYAN

    # ==========================================
    # SLIDE 4: WHAT (Key Results & Benchmarks)
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_bg(slide4)
    add_slide_header(slide4, "04", "RESULTS & BENCHMARKS", "WHAT We Delivered: Production Metrics", "Field-proven accuracy, high-throughput edge speed, and autonomous site operation.")

    # 2x2 Bento Metric Grid
    grid_w = Inches(5.72)
    grid_h = Inches(2.45)
    gx1 = Inches(0.8)
    gx2 = Inches(6.81)
    gy1 = Inches(1.75)
    gy2 = Inches(4.45)

    metrics_data = [
        # (x, y, badge_text, badge_color, stat_str, title_str, bullets)
        (gx1, gy1, "ACCURACY", C_CYAN, "94 – 97 %", "Correctly Identified", [
            "Consistent mask mAP across cartons & polybags",
            "High fidelity even on CPU-only Intel site PCs"
        ]),
        (gx2, gy1, "SPEED & LATENCY", C_BLUE, "~ 40 FPS", "Edge Processing Rate", [
            "Inference is significantly faster than the camera",
            "Camera stream framerate (~18 FPS) is the only ceiling"
        ]),
        (gx1, gy2, "DEPLOYED", C_EMERALD, "1 Site Live", "Industrial Production", [
            "One-click turnkey start & self-restart on boot",
            "30-day event history analytics & remote support"
        ]),
        (gx2, gy2, "DIAGNOSTICS", C_INDIGO, "7 Traffic Lights", "Real-Time Telemetry", [
            "Root-cause isolation in under 30 seconds",
            "Instant health checks — zero on-site visits needed"
        ]),
    ]

    for (mx, my, badge, bcolor, stat, label, bullets) in metrics_data:
        mcard = add_card(slide4, mx, my, grid_w, grid_h, bg_color=C_CARD_BG, border_color=C_CARD_BORDER)
        
        # Badge
        add_badge(slide4, mx + Inches(0.35), my + Inches(0.25), Inches(1.6), Inches(0.26), badge, RGBColor(20, 38, 68), bcolor, Pt(8.5))
        
        # Big Stat Text Box
        st_box = slide4.shapes.add_textbox(mx + Inches(0.35), my + Inches(0.6), Inches(2.6), Inches(0.8))
        tf_st = st_box.text_frame
        tf_st.word_wrap = True
        tf_st.margin_left = Inches(0)
        tf_st.margin_top = Inches(0)
        p_st = tf_st.paragraphs[0]
        r_st = p_st.add_run()
        r_st.text = stat
        r_st.font.name = FONT_HEADING
        r_st.font.size = Pt(26)
        r_st.font.bold = True
        r_st.font.color.rgb = C_TEXT_WHITE
        
        p_sub = tf_st.add_paragraph()
        r_sub = p_sub.add_run()
        r_sub.text = label
        r_sub.font.name = FONT_HEADING
        r_sub.font.size = Pt(11)
        r_sub.font.bold = True
        r_sub.font.color.rgb = bcolor

        # Bullets Text Box on Right side of card
        b_box = slide4.shapes.add_textbox(mx + Inches(2.9), my + Inches(0.4), grid_w - Inches(3.1), Inches(1.9))
        tf_b = b_box.text_frame
        tf_b.word_wrap = True
        tf_b.margin_left = Inches(0)
        
        for idx_b, bullet in enumerate(bullets):
            p_bl = tf_b.paragraphs[0] if idx_b == 0 else tf_b.add_paragraph()
            if idx_b > 0:
                p_bl.space_before = Pt(8)
            r_dot = p_bl.add_run()
            r_dot.text = "•  "
            r_dot.font.name = FONT_BODY
            r_dot.font.size = Pt(11)
            r_dot.font.bold = True
            r_dot.font.color.rgb = bcolor
            
            r_bt = p_bl.add_run()
            r_bt.text = bullet
            r_bt.font.name = FONT_BODY
            r_bt.font.size = Pt(11)
            r_bt.font.color.rgb = C_TEXT_SLATE

    # ==========================================
    # SLIDE 5: DEMO (Live Sorter Signal Verification)
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_bg(slide5)
    add_slide_header(slide5, "05", "LIVE VERIFICATION", "DEMO: Real-Time Sorter Output", "What the sorter PLC controller receives over UDP on every conveyor belt parcel crossing.")

    vid_left = Inches(0.8)
    vid_top = Inches(1.75)
    vid_w = Inches(7.5)
    vid_h = Inches(5.16)

    # Left Container Frame for Video
    vid_card = add_card(slide5, vid_left, vid_top, vid_w, vid_h, bg_color=C_CARD_BG, border_color=C_CARD_BORDER)
    add_badge(slide5, vid_left + Inches(0.3), vid_top + Inches(0.2), Inches(2.3), Inches(0.26), "LIVE FEED & DETECTION STREAM", RGBColor(20, 38, 68), C_CYAN, Pt(8.5))

    # Add the embedded movie
    movie_path = "/home/aatanda/logistic/.pptx_media/VAHK3IDLyaM.mp4"
    poster_path = "/home/aatanda/logistic/.pptx_media/image24.jpeg"
    if os.path.exists(movie_path):
        m_left = vid_left + Inches(0.3)
        m_top = vid_top + Inches(0.6)
        m_w = vid_w - Inches(0.6)
        m_h = vid_h - Inches(0.85)
        slide5.shapes.add_movie(
            movie_file=movie_path,
            left=m_left,
            top=m_top,
            width=m_w,
            height=m_h,
            poster_frame_image=poster_path if os.path.exists(poster_path) else None,
            mime_type="video/mp4"
        )

    # Right Column - 3 Telemetry / Protocol Specs Cards
    r_col_x = Inches(8.55)
    r_col_w = Inches(3.98)
    info_cards = [
        ("UDP TELEMETRY", "Port 9502 UDP Egress", "Transmits compact ~60-byte JSON packet per parcel crossing event directly to the PLC controller.", C_CYAN),
        ("RELIABILITY", "Single-Trigger Logic", "ByteTrack persistent ID guarantees each parcel fires exactly once at the virtual line without double-pulsing.", C_EMERALD),
        ("INTEGRATION", "Zero-Latency Routing", "Real-time classification gives the sorter maximum physical reaction window to divert parcels effortlessly.", C_BLUE),
    ]

    ic_h = Inches(1.58)
    ic_gap = Inches(0.2)
    ic_start_y = Inches(1.75)

    for i, (ibadge, ititle, idesc, icolor) in enumerate(info_cards):
        iy = ic_start_y + i * (ic_h + ic_gap)
        icard = add_card(slide5, r_col_x, iy, r_col_w, ic_h, bg_color=C_CARD_BG, border_color=C_CARD_BORDER)
        
        # Badge
        add_badge(slide5, r_col_x + Inches(0.25), iy + Inches(0.18), Inches(1.6), Inches(0.24), ibadge, RGBColor(20, 38, 68), icolor, Pt(8))
        
        # Text Frame
        tb_i = slide5.shapes.add_textbox(r_col_x + Inches(0.25), iy + Inches(0.48), r_col_w - Inches(0.5), ic_h - Inches(0.55))
        tf_i = tb_i.text_frame
        tf_i.word_wrap = True
        tf_i.margin_left = Inches(0)
        tf_i.margin_top = Inches(0)
        
        p_it = tf_i.paragraphs[0]
        r_it = p_it.add_run()
        r_it.text = ititle
        r_it.font.name = FONT_HEADING
        r_it.font.size = Pt(12)
        r_it.font.bold = True
        r_it.font.color.rgb = C_TEXT_WHITE
        
        p_id = tf_i.add_paragraph()
        p_id.space_before = Pt(3)
        r_id = p_id.add_run()
        r_id.text = idesc
        r_id.font.name = FONT_BODY
        r_id.font.size = Pt(10)
        r_id.font.color.rgb = C_TEXT_SLATE

    # Save to output path
    prs.save(output_path)
    print(f"Successfully created redesigned presentation at: {output_path}")

if __name__ == "__main__":
    create_deck("/home/aatanda/logistic/Project Celio v2.pptx")
